"""
Generate a sample Project Operations schedule template for local testing.

Mirrors the real template layout: Projects, Tasks, Assignments, Dependencies
and an Instructions sheet. Use this to exercise the parser and UI without the
sensitivity-labelled original.

    python make_sample.py
"""
from __future__ import annotations

from datetime import datetime

from openpyxl import Workbook


def build() -> Workbook:
    wb = Workbook()
    wb.properties.creator = "ScheduleAPI"
    wb.properties.lastModifiedBy = "ScheduleAPI"

    projects = wb.active
    projects.title = "Projects"
    projects.append(
           ["ProjectId", "ProjectName", "OwningCompany", "ContractingUnit",
            "CustomerAccount", "CalendarName", "ProjectManager", "ProjectStart"]
    )
    projects.append(
           ["PROJ-001", "Sample Migration Project A", "USPM", "US Consulting",
            "Contoso Ltd", "8 hr PST template", "SA Solutions Architect",
            datetime(2026, 7, 6, 9, 0)]
    )
    projects.append(
           ["PROJ-002", "Sample Rollout Project B", "USPM", "US Consulting",
            "Contoso Ltd", "8 hr PST template", "SA Solutions Architect",
            datetime(2026, 7, 13, 9, 0)]
    )

    tasks = wb.create_sheet("Tasks")
    tasks.append(
        ["ProjectId", "TaskKey", "TaskName", "ProjectBucket", "ParentTaskKey",
            "EffortHours", "ScheduledStart", "ScheduledEnd", "OutlineLevel",
            "TaskCol1"]
    )
    for r in [
        ["PROJ-001", "A-100", "Initiation", "Planning", None, 8, datetime(2026, 7, 6, 9), datetime(2026, 7, 6, 17), 1],
        ["PROJ-001", "A-110", "Confirm scope", "Planning", "A-100", 4, datetime(2026, 7, 6, 9), datetime(2026, 7, 6, 13), None],
        ["PROJ-001", "A-120", "Approve plan", "Planning", "A-100", 4, datetime(2026, 7, 6, 13), datetime(2026, 7, 6, 17), None],
        ["PROJ-001", "A-200", "Build", "Delivery", None, 24, datetime(2026, 7, 7, 9), datetime(2026, 7, 9, 17), 1],
        ["PROJ-001", "A-210", "Configure environment", "Delivery", "A-200", 16, datetime(2026, 7, 7, 9), datetime(2026, 7, 8, 17), None],
        ["PROJ-001", "A-220", "Validate data load", "Delivery", "A-200", 8, datetime(2026, 7, 9, 9), datetime(2026, 7, 9, 17), None],
        ["PROJ-002", "B-100", "Plan", "Planning", None, 8, datetime(2026, 7, 13, 9), datetime(2026, 7, 13, 17), 1],
        ["PROJ-002", "B-110", "Kickoff", "Planning", "B-100", 4, datetime(2026, 7, 13, 9), datetime(2026, 7, 13, 13), None],
        ["PROJ-002", "B-200", "Deploy", "Delivery", None, 16, datetime(2026, 7, 14, 9), datetime(2026, 7, 15, 17), 1],
    ]:
        tasks.append(r)

    assignments = wb.create_sheet("Assignments")
    assignments.append(
        ["ProjectId", "AssignmentKey", "TaskKey", "ResourceName", "RoleName"]
    )
    for r in [
        ["PROJ-001", "A-110-SA", "A-110", "SA Solutions Architect", "Consultant"],
        ["PROJ-001", "A-120-SA", "A-120", "SA Solutions Architect", "Consultant"],
        ["PROJ-001", "A-210-SA", "A-210", "SA Solutions Architect", "Consultant"],
        ["PROJ-001", "A-220-SA", "A-220", "SA Solutions Architect", "Consultant"],
        ["PROJ-002", "B-110-SA", "B-110", "SA Solutions Architect", "Consultant"],
        ["PROJ-002", "B-200-SA", "B-200", "SA Solutions Architect", "Consultant"],
    ]:
        assignments.append(r)

    deps = wb.create_sheet("Dependencies")
    deps.append(["ProjectId", "PredecessorTaskKey", "SuccessorTaskKey", "LinkType"])
    for r in [
        ["PROJ-001", "A-110", "A-120", "FinishToStart"],
        ["PROJ-001", "A-120", "A-210", "FinishToStart"],
        ["PROJ-001", "A-210", "A-220", "FinishToStart"],
        ["PROJ-002", "B-110", "B-200", "FinishToStart"],
    ]:
        deps.append(r)

    instr = wb.create_sheet("Instructions")
    instr.append(["Sheet", "Purpose", "Required keys"])
    instr.append(["Projects", "One row per project to create or reuse.", "ProjectId must be unique."])
    instr.append(["Tasks", "One row per task. Use ParentTaskKey for WBS hierarchy.", "ProjectId + TaskKey must be unique."])
    instr.append(["Assignments", "One row per task/resource assignment.", "ProjectId and TaskKey must exist."])
    instr.append(["Dependencies", "One row per task dependency.", "Predecessor/SuccessorTaskKey must exist in the same project."])

    return wb


if __name__ == "__main__":
    build().save("sample_template.xlsx")
    print("Wrote sample_template.xlsx")
